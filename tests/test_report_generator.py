"""
tests/test_report_generator.py — Модульные тесты генератора аналитических паспортов Excel (.xlsx).
"""
import io
import openpyxl
import pandas as pd
import pytest

import chemo_ml
import report_generator


def test_generate_excel_passport_plsda():
    """Тест генерации многостраничного паспорта для PLS-DA модели."""
    df, y = chemo_ml.generate_multimodal_benchmark()
    df["Class_Target"] = y.values
    X = df.drop(columns=["Sample_ID", "Class_Target"])

    res = chemo_ml.train_plsda_model(X, y.values, n_components=2)
    meta = {"project_name": "Baikal Test Benchmark", "operator": "Dr. Chemist"}

    excel_bytes = report_generator.generate_excel_passport(df, model_results=res, metadata=meta)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 5000

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Passport & Summary" in sheet_names
    assert "Fused Feature Matrix" in sheet_names
    assert "FT-ICR MS Descriptors" in sheet_names
    assert "EEM Descriptors" in sheet_names
    assert "UV-Vis Descriptors" in sheet_names
    assert "FTIR Descriptors" in sheet_names
    assert "Biomarkers & VIP" in sheet_names

    # Проверка листа сводки
    ws_sum = wb["Passport & Summary"]
    assert "CHEMOSUITE UNIFIED ANALYTICAL PASSPORT" in str(ws_sum["A1"].value)
    assert ws_sum["B5"].value == "Baikal Test Benchmark"


def test_generate_excel_passport_oplsda():
    """Тест генерации паспорта для OPLS-DA модели с вкладкой S-Plot."""
    df, y = chemo_ml.generate_multimodal_benchmark()
    df["Class_Target"] = y.values
    X = df.drop(columns=["Sample_ID", "Class_Target"])

    res = chemo_ml.train_oplsda_model(X, y.values, n_ortho=1)
    excel_bytes = report_generator.generate_excel_passport(df, model_results=res)

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "OPLS-DA S-Plot" in wb.sheetnames

    # Проверка содержимого S-Plot
    ws_s = wb["OPLS-DA S-Plot"]
    headers = [cell.value for cell in ws_s[1]]
    assert "Descriptor" in headers
    assert "p1_cov" in headers
    assert "p_corr" in headers
    assert "VIP" in headers


def test_generate_excel_passport_no_model():
    """Тест генерации паспорта без обученной модели (сырая матрица)."""
    df, y = chemo_ml.generate_multimodal_benchmark()
    df["Class_Target"] = y.values

    excel_bytes = report_generator.generate_excel_passport(df, model_results=None)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 3000

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Passport & Summary" in wb.sheetnames
    assert "Fused Feature Matrix" in wb.sheetnames
    assert "Biomarkers & VIP" not in wb.sheetnames


def test_generate_excel_passport_english():
    """Тест генерации многостраничного отчета на английском языке."""
    df, y = chemo_ml.generate_multimodal_benchmark()
    df["Class_Target"] = y.values
    X = df.drop(columns=["Sample_ID", "Class_Target"])

    res = chemo_ml.train_plsda_model(X, y.values, n_components=2)
    meta = {"project_name": "International Lake Monitoring", "operator": "Dr. Smith"}

    excel_bytes = report_generator.generate_excel_passport(df, model_results=res, metadata=meta, lang="en")
    assert isinstance(excel_bytes, bytes)

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws_sum = wb["Passport & Summary"]

    # Проверяем англоязычные заголовки
    assert ws_sum["A4"].value == "1. GENERAL ANALYSIS METADATA"
    assert ws_sum["A5"].value == "Project / Session:"
    assert ws_sum["A6"].value == "Generation Date:"
    assert ws_sum["A7"].value == "Operator / Engine:"
    assert ws_sum["A8"].value == "Total Samples:"
    assert ws_sum["B5"].value == "International Lake Monitoring"


def test_generate_excel_passport_elemental_and_crosscorr():
    """Тест генерации паспорта Excel с блоком Elemental (CHNS-O) и кросс-корреляциями."""
    df, y = chemo_ml.generate_multimodal_benchmark()
    df["Class_Target"] = y.values
    # Добавляем фиктивные дескрипторы CHNS-O
    df["Elem_H_C"] = 1.25
    df["Elem_O_C"] = 0.45
    df["Elem_N_C"] = 0.05

    cc_mock = {
        "top_cross_correlations": pd.DataFrame([
            {"Feature_A": "VK_1", "Block_A": "FT-ICR MS", "Feature_B": "C1_Humic", "Block_B": "EEM-PARAFAC", "Correlation": 0.85, "p_value": 0.001, "Significant": True}
        ])
    }

    excel_bytes = report_generator.generate_excel_passport(df, model_results=None, cross_corr_results=cc_mock)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))

    assert "CHNS-O Descriptors" in wb.sheetnames
    assert "Cross-Omics Correlations" in wb.sheetnames


def test_generate_ecological_passport_html_plsda():
    """Тест генерации интерактивного HTML экологического паспорта с моделью PLS-DA."""
    df, y = chemo_ml.generate_multimodal_benchmark()
    df["Class_Target"] = y.values
    X = df.drop(columns=["Sample_ID", "Class_Target"])

    res = chemo_ml.train_plsda_model(X, y.values, n_components=2)
    meta = {"project_name": "Lake Baikal Environmental Screening", "operator": "Dr. Chemist"}

    html_out = report_generator.generate_ecological_passport_html(df, model_results=res, metadata=meta, lang="ru")
    assert isinstance(html_out, str)
    assert "<!DOCTYPE html>" in html_out
    assert "Lake Baikal Environmental Screening" in html_out
    assert "PLS-DA" in html_out
    assert "R²X" in html_out
    assert "window.print()" in html_out
    assert "FT-ICR MS" in html_out
    assert "EEM-PARAFAC" in html_out
    assert "ЗАКЛЮЧЕНИЕ ЭКСПЕРТА:" in html_out


def test_generate_ecological_passport_html_english():
    """Тест генерации интерактивного HTML экологического паспорта на английском языке."""
    df, y = chemo_ml.generate_multimodal_benchmark()
    df["Class_Target"] = y.values

    html_out = report_generator.generate_ecological_passport_html(df, model_results=None, lang="en")
    assert isinstance(html_out, str)
    assert "<!DOCTYPE html>" in html_out
    assert "ChemoSuite Ecological Passport" in html_out
    assert "Print / Save as PDF" in html_out
    assert "EXPERT APPRAISAL:" in html_out


