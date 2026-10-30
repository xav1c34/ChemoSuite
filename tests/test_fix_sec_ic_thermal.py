"""
tests/test_fix_sec_ic_thermal.py — регрессионные тесты исправлений модулей SEC, IC, IRMS,
Thermal (TGA / Py-GC-MS), пакетного скрининга и генератора отчётов.

Каждый тест падает на прежней версии кода (неверный базис AsLS, линейные пресеты SEC,
подстрочное сопоставление столбцов IC, подстановка правдоподобных констант вместо NaN,
заготовленное «экспертное» заключение, отсутствие защиты от формульных инъекций и т.д.).
"""
import io
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

import batch_screener
import chemo_ml
import ic_core
import irms_core
import report_generator
import sec_core
import thermal_core

DEMO = Path(__file__).resolve().parent.parent / "demo_data"


# ==============================================================================
# SEC / GPC
# ==============================================================================

def test_sec_asls_flat_baseline_is_preserved():
    """(W + λD'D)z = Wy: плоская базовая линия 5.0 должна остаться 5.0 (раньше D·D' давал ~0 на краях)."""
    z = sec_core.baseline_als_sec(np.full(721, 5.0))
    assert np.allclose(z, 5.0, atol=1e-3)


@pytest.mark.parametrize("key", list(sec_core.CALIBRATION_PRESETS.keys()))
def test_sec_presets_reproduce_their_standards(key):
    """Коэффициенты пресета = МНК-подгонка его же стандартов; ошибка M ≤ 5 % (раньше до 38 %)."""
    preset = sec_core.CALIBRATION_PRESETS[key]
    v = np.array([s["vr"] for s in preset["standards"]])
    mw = np.array([s["mw"] for s in preset["standards"]])
    calib = sec_core.calibration_from_preset(key)
    rel_err = np.abs(np.asarray(calib.v_to_mw(v)) / mw - 1.0)
    assert rel_err.max() <= 0.05
    fresh = np.polyfit(v, np.log10(mw), deg=3)[::-1]
    assert np.allclose(preset["default_coeffs"], fresh, rtol=1e-6, atol=1e-9)
    assert calib.v_min == pytest.approx(v.min()) and calib.v_max == pytest.approx(v.max())


def test_sec_signal_outside_calibration_is_excluded_and_reported():
    """Сигнал вне диапазона стандартов не экстраполируется, а исключается с отчётом о доле площади."""
    v = np.linspace(4.0, 12.0, 721)
    calib = sec_core.calibration_from_preset("PSS_Sodium_Sulfonate")
    late = np.exp(-0.5 * ((v - 11.0) / 0.3) ** 2) * 100.0     # после легчайшего стандарта (10.25 мл)
    r_late = sec_core.calculate_sec_indices(v, late, calib)
    assert r_late["SEC_Below_Cal_Range_pct"] > 90.0
    assert r_late["SEC_Excluded_Area_pct"] == pytest.approx(r_late["SEC_Below_Cal_Range_pct"] + r_late["SEC_Above_Cal_Range_pct"], abs=0.05)
    centred = np.exp(-0.5 * ((v - 8.0) / 0.3) ** 2) * 100.0
    r_mid = sec_core.calculate_sec_indices(v, centred, calib)
    assert r_mid["SEC_Excluded_Area_pct"] < 0.1
    # Фракция IV (< 1 кДа) не покрыта калибровкой PSS (1.1–122 кДа) -> NaN, а не 0 %
    assert np.isnan(r_mid["SEC_Fraction_IV_pct"])


def test_sec_parse_accepts_user_calibration():
    """parse_sec_file принимает пользовательскую калибровку: log10 M = 7 − 0.4 V, пик при 8 мл -> Mp ≈ 6310 Да."""
    v = np.linspace(4.0, 12.0, 721)
    sig = np.exp(-0.5 * ((v - 8.0) / 0.3) ** 2) * 100.0
    txt = "Volume_mL,Signal_RI\n" + "\n".join(f"{a:.4f},{b:.5f}" for a, b in zip(v, sig))
    s = sec_core.parse_sec_file(txt, sample_id="X", calibration=[7.0, -0.4], calibration_v_range=(6.0, 10.0))
    assert s.calibration.calibration_type == "User_Coefficients"
    assert s.indices["SEC_Mp"] == pytest.approx(10 ** 3.8, rel=0.01)
    fitted = sec_core.fit_sec_calibration([6.0, 10.0], [10 ** 4.6, 10 ** 3.0], degree=1)
    s2 = sec_core.parse_sec_file(txt, sample_id="X", calibration=fitted)
    assert s2.indices["SEC_Mp"] == pytest.approx(10 ** 3.8, rel=0.01)


def test_sec_fit_requires_two_standards():
    """Раньше при < 2 стандартах молча подставлялся пресет PSS."""
    with pytest.raises(ValueError):
        sec_core.fit_sec_calibration([8.0], [5000.0])


# ==============================================================================
# Ion chromatography
# ==============================================================================

def test_ic_column_matching_is_exact_not_substring():
    """so4_mg_l -> SO4 (не Mg, по подстроке «mg»); sulfate не F; SO4_Cl_ratio не концентрация."""
    r = ic_core.resolve_ic_columns(["Sample_ID", "so4_mg_l", "SO4_Cl_ratio", "Mg", "F"])
    sp = r["species"]
    assert sp["so4_2minus"][0] == "so4_mg_l"
    assert sp["mg_2plus"][0] == "Mg"
    assert sp["f_minus"][0] == "F"
    assert "SO4_Cl_ratio" not in {v[0] for v in sp.values()}
    r2 = ic_core.resolve_ic_columns(["sulfate"])
    assert set(r2["species"]) == {"so4_2minus"}


def test_ic_baikal_demo_balances_with_bicarbonate():
    """Демо-набор Байкала (Falkner et al. 1991) с HCO3 ≈ 66 мг/л: |CBE| < 5 % (раньше −70…−90 %)."""
    samples = ic_core.parse_ic_table(ic_core.generate_synthetic_ic_dataset())
    baikal = [s for s in samples if "Baikal" in s.sample_id]
    assert baikal
    for s in baikal:
        assert abs(s.charge_balance_error_pct) < 5.0
        assert s.hco3_eff_mg_l == pytest.approx(66.0, abs=5.0)
        assert s.charge_balance_flag == "ok"


def test_ic_missing_bicarbonate_gives_nan_and_flag():
    """Без HCO3/щёлочности баланс зарядов не вычисляется (NaN + флаг), а не «−80 %»."""
    s = ic_core.ICSample("n", so4_2minus=5.5, cl_minus=0.4, na_plus=3.4, k_plus=1.0,
                         ca_2plus=16.0, mg_2plus=3.0).compute_metrics()
    assert np.isnan(s.charge_balance_error_pct)
    assert np.isnan(s.ionic_strength_mmol_l)
    assert s.charge_balance_flag == "HCO3_missing"


def test_ic_major_ion_missing_flag_and_so4_cl_nan():
    txt = "Sample_ID,SO4,Na,K,Ca,Mg,HCO3\nA,5.5,3.4,1.0,16,3.0,66\n"
    s = ic_core.parse_ic_table(txt)[0]
    assert s.charge_balance_flag == "major_ion_missing"
    assert "cl_minus" in s.metadata["missing_species"]
    assert np.isnan(s.so4_to_cl_ratio)          # Cl = 0 -> NaN, а не 5.5/1e-4


def test_ic_phosphate_charge_from_ph():
    """z̄(PO4) по pKa 2.15/7.20/12.35 (Stumm & Morgan 1996): pH 7.0 -> 1.387, pH 7.9 -> 1.834."""
    assert ic_core.phosphate_mean_charge(7.0)[0] == pytest.approx(1.387, abs=2e-3)
    assert ic_core.phosphate_mean_charge(7.9)[0] == pytest.approx(1.834, abs=2e-3)
    s = ic_core.ICSample("p", po4_3minus=1.0, so4_2minus=5.5, cl_minus=0.4, na_plus=3.4,
                         ca_2plus=16.0, mg_2plus=3.0, hco3_minus=66.0).compute_metrics()
    assert s.po4_mean_charge == 3.0 and any("PO4" in n for n in s.notes)


def test_ic_alkalinity_units():
    """Щёлочность 54.1 мг/л CaCO3 = 1.081 мэкв/л (SM 2320B, 50.04 мг CaCO3/мэкв); без единиц — игнор + предупреждение."""
    t = "Sample_ID,SO4,Cl,Na,K,Ca,Mg,Alkalinity (mg/L as CaCO3)\nA,5.5,0.4,3.4,1.0,16,3.0,54.1\n"
    a = ic_core.parse_ic_table(t)[0]
    assert a.alkalinity_meq_l == pytest.approx(54.1 / 50.04, rel=1e-3)
    assert a.carbonate_basis == "alkalinity"
    assert abs(a.charge_balance_error_pct) < 5.0
    t2 = "Sample_ID,SO4,Cl,Na,K,Ca,Mg,Alkalinity\nA,5.5,0.4,3.4,1.0,16,3.0,1.08\n"
    b = ic_core.parse_ic_table(t2)[0]
    assert b.alkalinity_meq_l is None
    assert any("Alkalinity" in w for w in b.metadata["parse_warnings"])
    b2 = ic_core.parse_ic_table(t2, alkalinity_unit="meq/L")[0]
    assert b2.alkalinity_meq_l == pytest.approx(1.08)


# ==============================================================================
# IRMS
# ==============================================================================

def test_irms_missing_isotope_is_nan_with_renormalised_probability():
    """Пропущенное δ13C -> NaN (раньше подставлялось −28 ‰); вероятность перенормирована, есть флаг."""
    txt = "Sample_ID,d13C,d15N,d34S,Sample_Type\nS1,,2.4,6.8,Sludge\n"
    s = irms_core.parse_irms_table(txt)[0]
    assert np.isnan(s.d13C)
    assert "d13C" not in s.kraft_prob_basis
    assert "kraft_prob_partial" in s.flags
    assert 0.0 <= s.kraft_lignin_prob <= 1.0


def test_irms_mixing_identical_endmembers_is_nan():
    """Совпадающие конечные члены -> NaN (раньше 0.5)."""
    assert np.isnan(irms_core.calculate_two_source_mixing(-25.0, -27.0, -27.0))
    f, flag = irms_core.calculate_two_source_mixing(-25.0, -27.0, -27.0, with_flag=True)
    assert np.isnan(f) and flag == "endmembers_identical"
    f2, flag2 = irms_core.calculate_two_source_mixing(-20.0, -27.0, -25.0, with_flag=True)
    assert f2 == 0.0 and flag2 == "outside_endmember_range"


def test_irms_d_excess_only_for_water():
    """d = δ2H − 8·δ18O (Dansgaard 1964) — только для воды."""
    w = irms_core.IRMSSample("W", d13C=float("nan"), d15N=float("nan"), d2H=-123.0, d18O=-15.8,
                             sample_type="Lake_Water").compute_metrics()
    assert w.deuterium_excess == pytest.approx(3.4)
    o = irms_core.IRMSSample("O", d13C=-27.0, d15N=2.0, d2H=-123.0, d18O=-15.8,
                             sample_type="Sludge").compute_metrics()
    assert np.isnan(o.deuterium_excess)


def test_irms_ratio_delta_and_two_point_normalisation():
    """R_VPDB = 0.011180; двухточечная нормализация (Paul et al. 2007)."""
    assert irms_core.ratio_to_delta(0.011180 * 1.01, "13C") == pytest.approx(10.0, abs=1e-9)
    assert irms_core.delta_to_ratio(10.0, "13C") == pytest.approx(0.011180 * 1.01, rel=1e-12)
    out = irms_core.two_point_normalization(np.array([-20.0, -30.0]), -10.5, -10.0, -40.2, -40.0)
    assert np.allclose(out, [-19.59596, -29.69697], atol=1e-4)
    assert np.isnan(irms_core.two_point_normalization(-20.0, -10.0, -10.0, -10.0, -40.0))


# ==============================================================================
# Thermal: Py-GC-MS
# ==============================================================================

@pytest.mark.parametrize("name,cls", [
    ("guaiacol", "G_Lignin"),
    ("Syringol", "S_Lignin"),
    ("2,6-dimethoxyphenol", "S_Lignin"),     # раньше «phenol» -> H_Lignin
    ("furfural", "Carbohydrate"),
    ("Сиреневая кислота", "S_Lignin"),       # раньше Unclassified
    ("Ацетованилон", "G_Lignin"),
])
def test_pyrolysis_name_classifier(name, cls):
    assert thermal_core.classify_pyrolysis_compound(name) == cls


def test_lignin_carbonyl_identification_precedence():
    """Ацетованилон — кетон, а не альдегид (подстрока «ванил»)."""
    assert thermal_core._identify_lignin_carbonyl("Ацетованилон") == "G_ketone"
    assert thermal_core._identify_lignin_carbonyl("Vanillic acid") == "G_acid"
    assert thermal_core._identify_lignin_carbonyl("Vanillin") == "G_aldehyde"


def _peak(name, area, cls="G_Lignin"):
    return thermal_core.PyGCMSPeak(retention_time=1.0, area=area, height=1.0, compound_name=name, chemical_class=cls)


def test_py_ac_al_is_acid_over_aldehyde_or_nan():
    """(Ac/Al)_G = кислота / альдегид (Hedges & Ertel 1982); без кислоты — NaN, а не 0.5."""
    idx = thermal_core.calculate_py_gcms_indices([_peak("Methyl vanillate", 30.0), _peak("Vanillin", 60.0)])
    assert idx["Py_Ac_Al_G"] == pytest.approx(0.5)
    assert idx["Py_Acid_Aldehyde_G"] == idx["Py_Ac_Al_G"]
    idx2 = thermal_core.calculate_py_gcms_indices([_peak("Vanillin", 60.0), _peak("Acetovanillone", 30.0)])
    assert np.isnan(idx2["Py_Ac_Al_G"])
    assert idx2["Py_Ketone_Aldehyde_G"] == pytest.approx(0.5)


def test_rt_match_score_has_no_floor():
    """RT_match_score = 1 − |ΔRT|/window без нижнего порога 0.5: ΔRT 0.45 при окне 0.5 -> 0.1."""
    m = thermal_core.match_py_compound(7.65, window=0.5)
    assert m["name_en"] == "Phenol"
    assert m["RT_match_score"] == pytest.approx(0.1, abs=1e-9)
    assert m["match_basis"] == "retention_time"


def test_named_peak_cas_comes_from_name_not_rt():
    """Пик «Syringol» на RT ванилина получает CAS сирингола, а не ванилина."""
    txt = "RT,Area,Name\n22.80,1000,Syringol\n"
    peak = thermal_core.parse_py_gcms_file(txt).peaks[0]
    assert peak.cas == "91-10-1"
    assert peak.chemical_class == "S_Lignin"
    assert peak.id_basis == "name" and np.isnan(peak.rt_match_score)


# ==============================================================================
# Thermal: TGA
# ==============================================================================

def test_tga_characteristic_temperatures_are_interpolated_or_nan():
    T = np.arange(25.0, 501.0, 50.0)
    m = 100.0 - 40.0 * (T - 25.0) / 475.0
    r = thermal_core.calculate_tga_indices(T, m)
    # m(150) = 89.47 %, T5 при m = 85.0 %: точно 203.1 °C (раньше — ближайший узел сетки 225 °C)
    assert r["TGA_T5_degC"] == pytest.approx(203.1, abs=0.1)
    assert np.isnan(r["TGA_T50_degC"])            # 50 % потери не достигнуто
    assert np.isnan(r["TGA_Char_Yield_800C"])     # 800 °C не достигнуто
    assert np.isnan(r["TGA_Residue_800C_pct"])


def test_tga_zone_iv_is_residue_at_final_temperature():
    T = np.linspace(25.0, 900.0, 200)
    m = 100.0 - 70.0 * (T - 25.0) / 875.0
    r = thermal_core.calculate_tga_indices(T, m)
    assert r["TGA_Zone_IV_pct"] == pytest.approx(30.0, abs=0.01)
    assert r["TGA_Residue_Final_pct"] == pytest.approx(30.0, abs=0.01)
    assert r["TGA_T_final_degC"] == pytest.approx(900.0)


def test_tga_dtg_keeps_negative_values():
    """Прирост массы (окисление) даёт DTG < 0 — раньше обрезалось до 0."""
    T = np.linspace(25.0, 900.0, 120)
    m = np.where(T < 600.0, 100.0 - 40.0 * (T - 25.0) / 575.0, 60.0 + 5.0 * (T - 600.0) / 300.0)
    txt = "Temperature (C),Weight (%)\n" + "\n".join(f"{a:.2f},{b:.4f}" for a, b in zip(T, m))
    s = thermal_core.parse_tga_file(txt, sample_id="ox")
    assert np.min(s.dtg) < 0.0


# ==============================================================================
# Batch screener
# ==============================================================================

_FOLDER_MODALITY = {
    "eem": "EEM-PARAFAC", "elemental": "Elemental", "fticr_ms": "FT-ICR MS", "ftir": "FTIR",
    "nmr": "NMR", "sec": "SEC/GPC", "thermal": "Thermal", "uv_vis": "UV-Vis",
}


def _demo_files():
    if not DEMO.exists():
        return []
    out = []
    for d in sorted(p for p in DEMO.iterdir() if p.is_dir()):
        for f in sorted(d.iterdir()):
            if f.suffix.lower() in (".csv", ".txt", ".tsv", ".jdx", ".dx"):
                out.append((d.name, f))
    return out


@pytest.mark.parametrize("neutral_name", [False, True])
def test_batch_detects_every_demo_file(neutral_name):
    files = _demo_files()
    if not files:
        pytest.skip("demo_data not found")
    bad = []
    for folder, f in files:
        name = ("data" + f.suffix) if neutral_name else f.name
        r = batch_screener.detect_file_modality(f.read_bytes(), name)
        if folder == "multimodal_ml":
            ok = r.get("kind") == "descriptor_matrix"
        else:
            ok = r["modality"] == _FOLDER_MODALITY[folder]
        if not ok:
            bad.append((str(f), r))
    assert not bad, bad


def test_batch_uv_header_is_not_elemental():
    body = b"wavelength,absorbance\n250,0.5\n260,0.45\n270,0.4\n280,0.38\n"
    assert batch_screener.detect_file_modality(body, "data.csv")["modality"] == "UV-Vis"


def test_batch_keeps_nan_for_missing_measurements():
    uv = DEMO / "uv_vis" / "Baikal_Control_01_uv.csv"
    ftir = DEMO / "ftir" / "Baikal_Pristine_01.csv"
    if not (uv.exists() and ftir.exists()):
        pytest.skip("demo_data not found")
    res = batch_screener.batch_screen_files([(uv.name, uv.read_bytes()), (ftir.name, ftir.read_bytes())])
    feats = res["fused_df"].drop(columns=["Sample_ID", "Class_Target"])
    assert len(res["fused_df"]) == 2
    assert feats.isna().any().any()               # раньше fillna(0.0)
    assert res["statistics"]["missing_values"] > 0


def test_batch_uv_spectrum_is_processed_without_assumed_doc():
    """UV-ветка падала (uv_samp.wavelengths); SUVA254 считалась с условным DOC = 5 мг/л."""
    uv = DEMO / "uv_vis" / "Baikal_Control_01_uv.csv"
    if not uv.exists():
        pytest.skip("demo_data not found")
    r = batch_screener.process_single_file(uv.read_bytes(), uv.name)
    assert r["status"] == "success", r["message"]
    assert "UV_A254" in r["descriptors"]
    assert "UV_SUVA254" not in r["descriptors"]          # нет измеренного DOC -> SUVA не выдумывается


def test_batch_fticr_uses_all_peaks_by_default():
    f = DEMO / "fticr_ms" / "Baikal_Control_01_fticr.csv"
    if not f.exists():
        pytest.skip("demo_data not found")
    r = batch_screener.process_single_file(f.read_bytes(), f.name)
    assert r["status"] == "success"
    assert "all 2200 peaks" in r["message"]
    r_top = batch_screener.process_single_file(f.read_bytes(), f.name, fticr_max_peaks=100)
    assert "top 100" in r_top["message"]


def test_batch_generic_csv_without_id_is_unrecognised():
    r = batch_screener.process_single_file(b"foo,bar\n1,2\n3,4\n", "data.csv")
    assert r["status"] == "warning"
    assert r["descriptors"] == {}
    assert "Unrecognised" in r["message"]


def test_batch_descriptor_matrix_rows_are_samples():
    f = DEMO / "multimodal_ml" / "chemo_unified_multimodal.csv"
    if not f.exists():
        pytest.skip("demo_data not found")
    r = batch_screener.process_single_file(f.read_bytes(), f.name)
    assert len(r["multi_samples"]) == len(pd.read_csv(f))


def test_batch_routes_pygcms_table_to_pyrolysis_indices():
    f = DEMO / "thermal" / "Baikal_NOM_01_PyGCMS.csv"
    if not f.exists():
        pytest.skip("demo_data not found")
    r = batch_screener.process_single_file(f.read_bytes(), "data.csv")
    assert r["modality"] == "Thermal"
    assert any(k.startswith("Py_") for k in r["descriptors"])


# ==============================================================================
# Report generator
# ==============================================================================

def _benchmark():
    df, y = chemo_ml.generate_multimodal_benchmark()
    df["Class_Target"] = y.values
    X = df.drop(columns=["Sample_ID", "Class_Target"])
    return df, chemo_ml.train_plsda_model(X, y.values, n_components=2)


def test_report_has_no_canned_lignin_claim_for_generic_classes():
    """Класс «Site_1» (содержит «1») раньше запускал вывод о шлам-лигнине БЦБК."""
    df = pd.DataFrame({
        "Sample_ID": [f"S{i}" for i in range(6)],
        "Class_Target": ["Site_1"] * 3 + ["Site_2"] * 3,
        "SUVA254": [2.1, 2.3, 2.2, 3.1, 3.0, 3.2],
        "E2_E3": [5.1, 5.0, 5.2, 4.1, 4.0, 4.2],
    })
    for lang in ("ru", "en"):
        h = report_generator.generate_ecological_passport_html(df, model_results=None, lang=lang)
        low = h.lower()
        assert "лигнин" not in low and "lignin" not in low
        assert "техногенная трансформация" not in low and "technogenic transformation" not in low
        assert "PASS (Chemometrically Validated)" not in h


def test_report_summary_is_data_driven():
    df, res = _benchmark()
    h = report_generator.generate_ecological_passport_html(df, model_results=res, lang="en")
    assert "Q² = " in h and "R²Y = " in h
    assert "Permutation test was not run" in h
    assert "UNVERIFIED" in h
    assert "small sample set (n = 14" in h
    assert "higher in class" in h
    assert "Largest block contributions" in h
    h2 = report_generator.generate_ecological_passport_html(df, model_results=res, lang="en",
                                                            permutation_results={"p_val_q2": 0.01})
    assert "Permutation test of Q²: p = 0.010" in h2
    assert "VALIDATED: Q²" in h2
    h3 = report_generator.generate_ecological_passport_html(df, model_results=res, lang="en",
                                                            permutation_results={"p_val_q2": 0.3})
    assert "NOT VALIDATED" in h3


def test_report_nan_metrics_are_not_printed_as_numbers():
    df, _ = _benchmark()
    res = {"model_type": "PLS-DA", "Q2": float("nan"), "R2Y": 80.0}
    h = report_generator.generate_ecological_passport_html(df, model_results=res, lang="en")
    assert "nan%" not in h
    xb = report_generator.generate_excel_passport(df, model_results=res, lang="en")
    ws = openpyxl.load_workbook(io.BytesIO(xb))["Passport & Summary"]
    vals = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}
    assert vals["R²X (Explained X-variance):"] == "—"      # раньше 0.00 %
    assert not any(isinstance(v, str) and "nan" in v.lower() for v in vals.values())


def test_excel_formula_injection_is_neutralised():
    df = pd.DataFrame({
        "Sample_ID": ['=HYPERLINK("http://evil.example","x")', "@SUM(A1)", "-28.5"],
        "Class_Target": ["A", "B", "A"],
        "=cmd": [1.0, 2.0, 3.0],
    })
    xb = report_generator.generate_excel_passport(df, model_results=None)
    ws = openpyxl.load_workbook(io.BytesIO(xb))["Fused Feature Matrix"]
    assert ws["A2"].value.startswith("'=HYPERLINK") and ws["A2"].data_type == "s"
    assert ws["A3"].value == "'@SUM(A1)"
    assert ws["A4"].value == "-28.5"                       # числовой литерал не портится
    assert ws["C1"].value == "'=cmd"
    safe = report_generator.sanitize_excel_df(df)
    assert safe["Sample_ID"].tolist() == ['\'=HYPERLINK("http://evil.example","x")', "'@SUM(A1)", "-28.5"]
    assert list(safe.columns)[2] == "'=cmd"
    assert df["Sample_ID"].iloc[0].startswith("=")         # исходный DataFrame не изменён


# ==============================================================================
# Py-GC-MS: пик-пикер TIC (базовая линия, порог по шуму, уникальные совпадения)
# ==============================================================================
def _tic_csv(rt, y):
    return pd.DataFrame({"Retention_Time_min": rt, "Intensity_TIC": y}).to_csv(index=False).encode()


def _g_and_s_refs():
    g = next(r for r in thermal_core.PY_GCMS_REFERENCE_LIBRARY if r["class"] == "G_Lignin")
    s = next(r for r in thermal_core.PY_GCMS_REFERENCE_LIBRARY if r["class"] == "S_Lignin")
    return g, s


@pytest.mark.parametrize("seed", range(3))
def test_py_tic_noise_bumps_are_not_lignin_peaks(seed):
    # базовая линия 50, шум σ=5, единственный пик — G-соединение; S-пиков нет
    g, _ = _g_and_s_refs()
    rt = np.arange(3.0, 35.0, 0.05)
    y = 50.0 + 250.0 * np.exp(-0.5 * ((rt - g["rt"]) / 0.06) ** 2) + np.random.default_rng(seed).normal(0, 5.0, rt.size)
    s = thermal_core.parse_py_gcms_file(_tic_csv(rt, y))
    # порог 3σ по высоте: в среднем ~1 ложный шумовой пик на хроматограмму, а не десятки
    assert len(s.peaks) <= 6
    assert not (s.indices["Py_S_pct"] > 10.0)
    named = [p.name_en for p in s.peaks if p.id_basis == "retention_time"]
    assert len(named) == len(set(named))          # одно соединение — один пик
    g_peaks = [p for p in s.peaks if p.name_en == g["name_en"]]
    assert len(g_peaks) == 1
    true_area = 250.0 * 0.06 * np.sqrt(2.0 * np.pi)   # площадь над базовой линией, а не «пол × окно»
    assert g_peaks[0].area == pytest.approx(true_area, rel=0.15)


def test_py_tic_recovers_g_to_s_ratio():
    g, s_ref = _g_and_s_refs()
    rt = np.arange(3.0, 35.0, 0.05)
    y = (50.0 + 300.0 * np.exp(-0.5 * ((rt - g["rt"]) / 0.06) ** 2)
         + 60.0 * np.exp(-0.5 * ((rt - s_ref["rt"]) / 0.06) ** 2)
         + np.random.default_rng(7).normal(0, 3.0, rt.size))
    s = thermal_core.parse_py_gcms_file(_tic_csv(rt, y))
    assert s.indices["Py_S_G_ratio"] == pytest.approx(60.0 / 300.0, rel=0.2)
